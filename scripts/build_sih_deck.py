import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

src_file = "d:/project/Crimenet/CN v1.pptx"
dst_file = "d:/project/Crimenet/SIH2026_VEILLE_Official_Format.pptx"

prs = Presentation(src_file)

# Text replacements across all slides to clean up typos / old encoding
replacements = {
    "Problem Statement ID ? SIH26189": "Problem Statement ID: SIH26189",
    "Team Name ? Team Void": "Team Name: Team Void",
    " ? Innovation": " - Innovation",
    "?risk scores?": "'risk scores'",
    "?Graph-DB + LLM?": "'Graph-DB + LLM'",
    "?candidate?": "'candidate'",
    "? defensible for": "- defensible for",
    "? confirming": "- confirming",
    "? Neo4j": "- Neo4j",
    "Postgres ? Neo4j": "Postgres -> Neo4j",
    "stack ? FastAPI": "stack - FastAPI",
    "mission ? 15,000+": "mission - 15,000+",
    "lock-in.": "lock-in.",
    "scores ? human": "scores - human",
    "burden.": "burden.",
    "demo.": "demo.",
    "? but are complex": "- but are complex",
    "?  IBM": "-  IBM",
    "?  Generic": "-  Generic",
    "Palantir Gotham  ?  IBM": "Palantir Gotham  |  IBM",
    "Notebook  ?  Generic": "Notebook  |  Generic",
    "scores ? human review gates every decision": "scores - human review gates every decision"
}

for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    for old_str, new_str in replacements.items():
                        if old_str in r.text:
                            r.text = r.text.replace(old_str, new_str)
                    r.text = r.text.replace(" ? ", " - ")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                if child.has_text_frame:
                    for p in child.text_frame.paragraphs:
                        for r in p.runs:
                            r.text = r.text.replace(" ? ", " - ")

# Update Slide 2: Replace the workflow image with the new crisp slide2_workflow.png
slide2 = prs.slides[1]
for s in list(slide2.shapes):
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP and "Group 54" in s.name:
        # Get position of Group 54
        left, top, width, height = s.left, s.top, s.width, s.height
        # Remove old group shape
        sp_elem = s._element
        sp_elem.getparent().remove(sp_elem)
        # Add new high-res workflow flowchart
        slide2.shapes.add_picture('d:/project/Crimenet/diagrams_v2/slide2_workflow.png',
                                  left, top, width, height)
        break

# Update Slide 3: Update Picture 16 with new crisp slide3_swimlane.png
slide3 = prs.slides[2]
for s in list(slide3.shapes):
    if s.name == "Picture 16":
        left, top, width, height = s.left, s.top, s.width, s.height
        sp_elem = s._element
        sp_elem.getparent().remove(sp_elem)
        slide3.shapes.add_picture('d:/project/Crimenet/diagrams_v2/slide3_swimlane.png',
                                  left, top, width, height)
        break

prs.save(dst_file)
print(f"Perfected presentation saved to: {dst_file}")
